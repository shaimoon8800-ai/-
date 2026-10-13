#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""数据体检：给一张 CSV/Excel 出字段画像与问题清单（不改动原文件）。

依赖: openpyxl（仅读 .xlsx 时需要）；CSV 走标准库。

用法:
  python3 profile_data.py --input 数据.csv --out 体检报告.md
  python3 profile_data.py --input 数据.xlsx --json
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

CN = re.compile(r"[\u4e00-\u9fa5]")
NUM_LIKE = re.compile(r"^[\s¥￥$]*-?[\d,]+(\.\d+)?\s*[%元万元]?$")
DATE_PATS = [
    re.compile(r"^\d{4}[-/.]\d{1,2}[-/.]\d{1,2}$"),
    re.compile(r"^\d{8}$"),
    re.compile(r"^\d{4}年\d{1,2}月\d{1,2}日$"),
]
FULLWIDTH = {chr(0xFF01 + i): chr(0x21 + i) for i in range(94)}


def read_table(path: Path) -> tuple[list[str], list[list[str]]]:
    ext = path.suffix.lower()
    if ext in (".csv", ".tsv", ".txt"):
        delim = "\t" if ext == ".tsv" else ","
        for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
            try:
                with path.open(encoding=enc, newline="") as f:
                    rows = list(csv.reader(f, delimiter=delim))
                break
            except UnicodeDecodeError:
                continue
        else:
            raise SystemExit("[FATAL] 无法解码 CSV，请确认编码")
        if not rows:
            raise SystemExit("[FATAL] 文件为空")
        return rows[0], rows[1:]

    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            raise SystemExit("[FATAL] 读 xlsx 需要 openpyxl：pip install openpyxl")
        wb = load_workbook(str(path), read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        if not rows:
            raise SystemExit("[FATAL] 表为空")
        return rows[0], rows[1:]

    raise SystemExit(f"[FATAL] 不支持的格式: {ext}")


def is_blank(v: str) -> bool:
    return v is None or str(v).strip() == ""


def norm(s: str) -> str:
    """归一化用于类型判断：全角转半角 + strip。"""
    return "".join(FULLWIDTH.get(ch, ch) for ch in str(s)).strip()


def classify(values: list[str]) -> str:
    vals = [v for v in values if not is_blank(v)]
    if not vals:
        return "空列"
    n_num = n_date = n_cn = 0
    for v in vals:
        v2 = norm(v)
        if any(p.match(v2) for p in DATE_PATS):
            n_date += 1
        elif NUM_LIKE.match(v2):
            n_num += 1
        if CN.search(v2):
            n_cn += 1
    t = len(vals)
    if n_date / t > 0.8:
        return "日期"
    if n_num / t > 0.8:
        return "数值(含格式)" if any(("," in norm(v) or "¥" in v or "元" in v) for v in vals) else "数值"
    if n_cn / t > 0.5:
        return "文本"
    return "混合"


def profile_col(name: str, values: list[str], total: int) -> dict:
    nonblank = [v for v in values if not is_blank(v)]
    missing = total - len(nonblank)
    uniq = len(set(norm(v) for v in nonblank))
    kind = classify(values)
    issues: list[str] = []

    if missing:
        rate = missing * 100 / max(total, 1)
        lvl = "🔴" if rate > 50 else ("🟡" if rate > 5 else "🟢")
        issues.append(f"{lvl} 缺失 {missing} 行（{rate:.1f}%）")

    if nonblank:
        n_space = sum(1 for v in nonblank if str(v) != str(v).strip())
        if n_space:
            issues.append(f"🟢 首尾空格 {n_space} 行")
        n_full = sum(1 for v in nonblank if any(ord(c) > 0xFEE0 for c in str(v)))
        if n_full:
            issues.append(f"🟡 含全角字符 {n_full} 行")
        if kind.startswith("数值"):
            bad = [v for v in nonblank if not NUM_LIKE.match(norm(v))]
            if bad:
                issues.append(f"🟡 数值列混入非数值 {len(bad)} 行，例：{bad[0][:20]!r}")
            nums = []
            for v in nonblank:
                s = re.sub(r"[¥￥$,元万元%\s]", "", norm(v))
                try:
                    nums.append(float(s))
                except ValueError:
                    pass
            if nums:
                neg = [n for n in nums if n < 0]
                if neg:
                    issues.append(f"🟡 存在负值 {len(neg)} 个（确认是否退款/冲销）")
                if len(nums) > 4:
                    # 箱线图 IQR 法（比「均值×N」稳，均值会被极端值本身拉高而失效）
                    s = sorted(nums)
                    med = s[len(s) // 2]
                    q1 = s[len(s) // 4]
                    q3 = s[(len(s) * 3) // 4]
                    iqr = q3 - q1
                    if iqr > 0:
                        lo, hi = q1 - 1.5 * iqr, q3 + 1.5 * iqr
                        big = [n for n in nums if n < lo or n > hi]
                    else:
                        big = [n for n in nums if med and abs(n) > abs(med) * 50]
                    if big:
                        issues.append(
                            f"🟡 疑似极端值 {len(big)} 个（IQR 法），"
                            f"最大 {max(big):,.2f}，最小 {min(big):,.2f}")
        if kind == "日期":
            forms = set()
            for v in nonblank:
                v2 = norm(v)
                if re.match(r"^\d{8}$", v2):
                    forms.add("YYYYMMDD")
                elif "/" in v2:
                    forms.add("YYYY/M/D")
                elif "." in v2:
                    forms.add("YYYY.M.D")
                elif "年" in v2:
                    forms.add("YYYY年M月D日")
                elif "-" in v2:
                    forms.add("YYYY-MM-DD")
            if len(forms) > 1:
                issues.append(f"🟡 日期格式不统一：{', '.join(sorted(forms))}")

    return {
        "name": name,
        "type": kind,
        "missing": missing,
        "missing_rate": round(missing * 100 / max(total, 1), 1),
        "unique": uniq,
        "is_key": (missing == 0 and uniq == total and total > 0),
        "sample": [str(v)[:18] for v in nonblank[:3]],
        "issues": issues,
    }


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--out", default="")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args()

    p = Path(args.input)
    if not p.exists():
        print(f"[FATAL] 文件不存在: {p}")
        return 2

    header, rows = read_table(p)
    ncol = len(header)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    total = len(rows)

    # 完全重复行
    seen: Counter = Counter(tuple(norm(c) for c in r) for r in rows)
    dup = sum(v - 1 for v in seen.values() if v > 1)

    profiles = [profile_col(header[i] or f"列{i+1}",
                            [r[i] for r in rows], total) for i in range(ncol)]
    keys = [c["name"] for c in profiles if c["is_key"]]

    all_issues = [f"{c['name']}: {iss}" for c in profiles for iss in c["issues"]]
    blocking = [i for i in all_issues if "🔴" in i]
    affecting = [i for i in all_issues if "🟡" in i]

    result = {
        "file": p.name,
        "rows": total,
        "cols": ncol,
        "duplicate_rows": dup,
        "candidate_keys": keys,
        "columns": profiles,
        "blocking": blocking,
        "affecting": affecting,
    }

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    L = [
        f"# 数据体检报告：{p.name}", "",
        f"- 行数：**{total}**　列数：**{ncol}**",
        f"- 完全重复行：**{dup}**",
        f"- 候选主键：{('、'.join(keys)) if keys else '**无**（所有列都有重复或缺失）'}",
        "",
        "## 字段画像", "",
        "| 列 | 类型 | 缺失 | 缺失率 | 唯一值 | 样例 |",
        "|---|---|---|---|---|---|",
    ]
    for c in profiles:
        L.append(f"| {c['name']} | {c['type']} | {c['missing']} | {c['missing_rate']}% "
                 f"| {c['unique']} | {' / '.join(c['sample'])[:40]} |")

    L += ["", "## 问题清单", ""]
    if blocking:
        L += ["### 🔴 阻断级（必须先解决）", ""] + [f"- {i}" for i in blocking] + [""]
    if affecting:
        L += ["### 🟡 影响分析", ""] + [f"- {i}" for i in affecting] + [""]
    minor = [i for i in all_issues if i not in blocking and i not in affecting]
    if minor:
        L += ["### 🟢 美化级", ""] + [f"- {i}" for i in minor] + [""]

    if not all_issues:
        L += ["✅ 未发现明显问题，可以直接分析。", ""]

    L += ["", "## 建议", ""]
    if blocking:
        L += ["1. 先解决 🔴 问题，否则后续分析结论不可靠"]
    if dup:
        L += [f"2. 存在 {dup} 行完全重复，建议去重（需确认是否为同一笔记录）"]
    if affecting:
        L += ["3. 清洗前先备份原文件；清洗后逐项核对日志"]
    if not blocking and not affecting:
        L += ["数据质量良好，可直接进入分析阶段（透视表 / 图表）"]

    report = "\n".join(L)
    if args.out:
        Path(args.out).write_text(report, encoding="utf-8")
        print(f"[OK] 体检报告已写出：{args.out}")
    print(report)
    return 0


if __name__ == "__main__":
    sys.exit(main())

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把提取结果生成规范的 Excel 报销台账（三个 sheet：明细 / 汇总 / 异常）。

依赖: openpyxl
用法:
  python3 build_ledger.py --input records.json --out 报销台账.xlsx
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    print("[FATAL] 需要 openpyxl：pip install openpyxl")
    raise SystemExit(2)

HEADER_FILL = PatternFill("solid", fgColor="F2F3F5")
HEADER_FONT = Font(bold=True, size=11)
RED_FILL = PatternFill("solid", fgColor="FDE8E8")
YELLOW_FILL = PatternFill("solid", fgColor="FEF3C7")
THIN = Side(style="thin", color="D9DCE1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

DETAIL_COLS = [
    ("序号", 6), ("开票日期", 12), ("发票号码", 24), ("销售方", 30),
    ("金额(元)", 12), ("税额(元)", 11), ("类别", 10), ("票种", 13),
    ("状态", 10), ("来源文件", 30), ("备注", 30),
]


def dedupe(records: list[dict]) -> list[dict]:
    """按发票号去重；同号保留首张，其余标『重复』。再标『疑似重复』。"""
    seen_no: dict[str, int] = {}
    seen_fuzzy: dict[tuple, int] = {}
    for i, r in enumerate(records, 1):
        r.setdefault("_notes", [])
        no = r.get("invoice_no")
        if no:
            if no in seen_no:
                r["status"] = "重复"
                r["_notes"].append(f"发票号与第 {seen_no[no]} 行重复")
                continue
            seen_no[no] = i
        key = (r.get("date"), r.get("amount"), r.get("seller"))
        if all(key) and key in seen_fuzzy:
            r["_notes"].append(f"疑似与第 {seen_fuzzy[key]} 行重复（同日期/金额/开票方）")
            if r.get("status") == "有效":
                r["status"] = "待确认"
        elif all(key):
            seen_fuzzy[key] = i
    return records


def build(records: list[dict], out_path: Path) -> tuple[dict, list[dict]]:
    wb = Workbook()

    # ---------- 明细 ----------
    ws = wb.active
    ws.title = "明细"
    for ci, (name, width) in enumerate(DETAIL_COLS, 1):
        c = ws.cell(row=1, column=ci, value=name)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.freeze_panes = "A2"

    valid = [r for r in records if r.get("status") != "重复"]
    for idx, r in enumerate(valid, 1):
        row = idx + 1
        ws.cell(row=row, column=1, value=idx)
        ws.cell(row=row, column=2, value=r.get("date") or "")
        ws.cell(row=row, column=3, value=r.get("invoice_no") or "")
        ws.cell(row=row, column=4, value=r.get("seller") or "")
        amt = ws.cell(row=row, column=5, value=r.get("amount"))
        amt.number_format = "#,##0.00"
        tax = ws.cell(row=row, column=6, value=r.get("tax"))
        tax.number_format = "#,##0.00"
        ws.cell(row=row, column=7, value=r.get("category") or "")
        ws.cell(row=row, column=8, value=r.get("type") or "")
        sc = ws.cell(row=row, column=9, value=r.get("status") or "")
        ws.cell(row=row, column=10, value=r.get("file") or "")
        ws.cell(row=row, column=11, value="；".join(r.get("_notes", [])) or "")
        for ci in range(1, len(DETAIL_COLS) + 1):
            ws.cell(row=row, column=ci).border = BORDER
        if r.get("status") == "重复":
            sc.fill = RED_FILL
        elif r.get("status") in ("待确认", "待人工补录"):
            sc.fill = YELLOW_FILL

    # 合计行
    total_row = len(valid) + 2
    ws.cell(row=total_row, column=4, value="合计").font = Font(bold=True)
    sum_cell = ws.cell(row=total_row, column=5,
                       value=f"=SUM(E2:E{len(valid)+1})" if valid else 0)
    sum_cell.number_format = "#,##0.00"
    sum_cell.font = Font(bold=True)
    tax_sum = ws.cell(row=total_row, column=6,
                      value=f"=SUM(F2:F{len(valid)+1})" if valid else 0)
    tax_sum.number_format = "#,##0.00"
    tax_sum.font = Font(bold=True)

    # ---------- 汇总 ----------
    ws2 = wb.create_sheet("汇总")
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 10
    ws2.column_dimensions["C"].width = 14
    ws2.column_dimensions["D"].width = 10
    row = 1

    def block(title: str, headers: list[str], rows: list[list]) -> int:
        nonlocal row
        ws2.cell(row=row, column=1, value=title).font = Font(bold=True, size=12)
        row += 1
        for ci, h in enumerate(headers, 1):
            c = ws2.cell(row=row, column=ci, value=h)
            c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        row += 1
        for r_ in rows:
            for ci, v in enumerate(r_, 1):
                c = ws2.cell(row=row, column=ci, value=v)
                c.border = BORDER
                if ci == 3 and isinstance(v, (int, float)):
                    c.number_format = "#,##0.00"
            row += 1
        row += 1
        return row

    by_month: dict[str, list[float]] = defaultdict(list)
    by_cat: dict[str, list[float]] = defaultdict(list)
    by_seller: dict[str, list[float]] = defaultdict(list)
    for r in valid:
        amt = r.get("amount") or 0
        month = (r.get("date") or "未知")[:7]
        by_month[month].append(amt)
        by_cat[r.get("category") or "其他"].append(amt)
        by_seller[r.get("seller") or "未知"].append(amt)

    grand = sum(sum(v) for v in by_month.values())
    block("① 按月份",
          ["月份", "张数", "金额(元)"],
          [[m, len(v), round(sum(v), 2)] for m, v in sorted(by_month.items())])

    def pct(v: float) -> str:
        return f"{v*100/grand:.1f}%" if grand else "0%"

    block("② 按类别",
          ["类别", "张数", "金额(元)", "占比"],
          [[k, len(v), round(sum(v), 2), pct(sum(v))]
           for k, v in sorted(by_cat.items(), key=lambda x: -sum(x[1]))])

    top_sellers = sorted(by_seller.items(), key=lambda x: -sum(x[1]))[:10]
    block("③ 开票方 Top 10",
          ["销售方", "张数", "金额(元)"],
          [[k, len(v), round(sum(v), 2)] for k, v in top_sellers])

    ws2.cell(row=row, column=1, value="合计").font = Font(bold=True)
    c = ws2.cell(row=row, column=3, value=round(grand, 2))
    c.number_format = "#,##0.00"
    c.font = Font(bold=True)

    # ---------- 异常 ----------
    ws3 = wb.create_sheet("异常")
    for ci, (name, width) in enumerate(
            [("类型", 14), ("发票号码", 24), ("来源文件", 30), ("问题描述", 46), ("建议处置", 20)], 1):
        c = ws3.cell(row=1, column=ci, value=name)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        ws3.column_dimensions[get_column_letter(ci)].width = width
    ws3.freeze_panes = "A2"

    anomalies: list[dict] = []
    r3 = 2
    for rec in records:
        problems = list(rec.get("warnings", []))
        if rec.get("status") == "重复":
            problems.insert(0, "发票号码重复")
        if rec.get("status") == "待人工补录":
            problems.insert(0, "无法自动提取（扫描件/图片）")
        if rec.get("amount") in (0, None):
            problems.append("金额缺失或为 0")
        if not problems:
            continue
        kind = ("重复" if rec.get("status") == "重复"
                else "待补录" if rec.get("status") == "待人工补录" else "缺字段")
        advice = ("人工确认后删除重复行" if kind == "重复"
                  else "按模板补录后重新导入" if kind == "待补录"
                  else "核对原始票据并补齐")
        ws3.cell(row=r3, column=1, value=kind)
        ws3.cell(row=r3, column=2, value=rec.get("invoice_no") or "—")
        ws3.cell(row=r3, column=3, value=rec.get("file") or "")
        ws3.cell(row=r3, column=4, value="；".join(problems))
        ws3.cell(row=r3, column=5, value=advice)
        for ci in range(1, 6):
            ws3.cell(row=r3, column=ci).border = BORDER
        anomalies.append({"类型": kind, "文件": rec.get("file"),
                          "问题": "；".join(problems)})
        r3 += 1

    if r3 == 2:
        ws3.cell(row=2, column=1, value="（无异常）")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))

    summary = {
        "total": len(records),
        "valid": len(valid),
        "duplicate": sum(1 for r in records if r.get("status") == "重复"),
        "manual": sum(1 for r in records if r.get("status") == "待人工补录"),
        "amount": round(grand, 2),
        "by_category": {k: round(sum(v), 2) for k, v in by_cat.items()},
        "by_month": {k: round(sum(v), 2) for k, v in by_month.items()},
        "anomalies": anomalies,
    }
    return summary, anomalies


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="extract_invoices.py 产出的 records.json")
    ap.add_argument("--out", default="报销台账.xlsx")
    ap.add_argument("--summary-json", default="", help="同时输出汇总 JSON")
    args = ap.parse_args()

    p = Path(args.input)
    if not p.exists():
        print(f"[FATAL] 找不到 {p}")
        return 2
    data = json.loads(p.read_text(encoding="utf-8"))
    records = data.get("records", [])
    if not records:
        print("[FATAL] records.json 里没有记录")
        return 2

    records = dedupe(records)
    out = Path(args.out)
    summary, anomalies = build(records, out)

    print(f"[OK] 台账已生成: {out}")
    print(f"     总票据 {summary['total']}　有效 {summary['valid']}　"
          f"重复 {summary['duplicate']}　待补录 {summary['manual']}")
    print(f"     合计金额 ¥{summary['amount']:,.2f}")
    if summary["by_category"]:
        print("     按类别：")
        for k, v in sorted(summary["by_category"].items(), key=lambda x: -x[1]):
            print(f"       - {k}: ¥{v:,.2f}")
    if anomalies:
        print(f"     ⚠️ 异常 {len(anomalies)} 项，见『异常』sheet")

    if args.summary_json:
        Path(args.summary_json).write_text(
            json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[OK] 汇总 JSON: {args.summary_json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
